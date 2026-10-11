"""Rasm → Word (Gemini vision). `fix` rejimida imlo xatolari tuzatiladi va formatlash tiklanadi."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from ..ai.clients import gemini_json

PROMPT = (
    "Rasmdagi hujjat matnini to'liq va aynan o'qib ol. Tuzilishini saqla: sarlavhalar, abzatslar, ro'yxatlar, "
    "jadvallar. {fix}\n"
    'Faqat JSON: {{"blocks": [{{"type": "heading|paragraph|list|table", "text": "...", '
    '"align": "left|center|right|justify", "bold": false, "items": ["ro\'yxat uchun"], '
    '"rows": [["jadval", "uchun"]]}}]}}'
)
FIX = ("Imlo va tinish belgisi xatolarini tuzat, so'zlar bo'linib qolgan bo'lsa birlashtir, "
       "xira joylarni mazmunga qarab tikla. Ma'noni o'zgartirma.")
NOFIX = "Hech narsani o'zgartirma, xatolar bo'lsa ham aynan ko'chir."

ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}


async def recognize(image: bytes, mime: str, fix: bool) -> list[dict]:
    d = await gemini_json(PROMPT.format(fix=FIX if fix else NOFIX), [(image, mime)], temperature=0.1)
    return d.get("blocks", [])


def build_docx(pages: list[list[dict]], out: Path, fix: bool) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3 if fix else 2), Cm(1.5 if fix else 2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(14 if fix else 12)
    if fix:
        st.paragraph_format.line_spacing = 1.5
    for pi, blocks in enumerate(pages):
        if pi:
            doc.add_page_break()
        for b in blocks:
            kind = b.get("type", "paragraph")
            if kind == "table" and b.get("rows"):
                rows = b["rows"]
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncols)
                t.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j, c in enumerate(r):
                        t.cell(i, j).text = str(c)
                continue
            if kind == "list":
                for it in b.get("items") or [b.get("text", "")]:
                    doc.add_paragraph(str(it), style="List Bullet")
                continue
            text = b.get("text", "")
            if not text:
                continue
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.bold = bool(b.get("bold")) or kind == "heading"
            if kind == "heading":
                r.font.size = Pt(16)
            p.alignment = ALIGN.get(b.get("align", ""), WD_ALIGN_PARAGRAPH.CENTER if kind == "heading"
                                    else (WD_ALIGN_PARAGRAPH.JUSTIFY if fix else WD_ALIGN_PARAGRAPH.LEFT))
            if fix and kind == "paragraph":
                p.paragraph_format.first_line_indent = Cm(1.25)
    doc.save(out)
    return out
