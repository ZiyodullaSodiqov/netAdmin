"""Test savollari: til, fan, mavzu, sinf asosida — variantlarga ajratilgan DOCX + javoblar kaliti."""
from __future__ import annotations

import asyncio
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..ai.clients import openai_json
from ..services import LANG_FULL
from .docs import _pborder, _shade

LABELS = {
    "uz": {"variant": "{n}-variant", "class": "{n}-sinf", "keys": "JAVOBLAR KALITI", "name": "F.I.Sh: ____________",
           "subject": "Fan", "topic": "Mavzu", "date": "Sana: ____________", "question": "Savol",
           "answer": "Javob"},
    "ru": {"variant": "Вариант {n}", "class": "{n} класс", "keys": "КЛЮЧИ ОТВЕТОВ", "name": "Ф.И.О.: ____________",
           "subject": "Предмет", "topic": "Тема", "date": "Дата: ____________", "question": "Вопрос",
           "answer": "Ответ"},
    "en": {"variant": "Variant {n}", "class": "Grade {n}", "keys": "ANSWER KEY", "name": "Name: ____________",
           "subject": "Subject", "topic": "Topic", "date": "Date: ____________", "question": "Question",
           "answer": "Answer"},
}
LETTERS = "ABCD"

# 5 uslub — tuzilishi har xil:
# list — klassik ro'yxat; grid — rangli sarlavha lentasi, javoblar 2x2; cards — har savol ramkali kartochkada,
# yonida «Javob» katagi; table — barcha savollar bitta jadvalda (A/B/C/D ustunlari); columns — ikki ustunli sahifa
STYLES = [
    {"kind": "list", "font": "Times New Roman", "head": None},
    {"kind": "grid", "font": "Arial", "head": "1F4E79"},
    {"kind": "cards", "font": "Calibri", "head": "0F6B57"},
    {"kind": "table", "font": "Cambria", "head": "7A1F3D"},
    {"kind": "columns", "font": "Georgia", "head": "5B3FA8"},
]


async def generate_variants(subject: str, topic: str, grade: str, lang: str, variants: int, count: int) -> list[list[dict]]:
    system = (f"Sen {grade}-sinf o'quvchilari uchun test tuzuvchi tajribali o'qituvchisan. Faqat JSON qaytar. "
              f"Barcha savol va javoblar {LANG_FULL[lang]} tilida.")

    async def one(n: int) -> list[dict]:
        d = await openai_json(system, (
            f"Fan: {subject}. Mavzu: {topic}. Sinf: {grade}. Bu {n}-variant (boshqa variantlardan farqli savollar).\n"
            f"Aynan {count} ta test savoli tuz, har birida 4 ta javob varianti, faqat bittasi to'g'ri. "
            "Qiyinlik sinf darajasiga mos, osondan qiyinga. To'g'ri javob o'rni tasodifiy bo'lsin.\n"
            'Format: {"questions": [{"q": "...", "options": ["..","..","..",".."], "answer": 0}]} '
            "(answer — to'g'ri javob indeksi 0-3)"
        ), temperature=0.8)
        qs = []
        for q in d.get("questions", []):
            opts = [str(o) for o in q.get("options", [])][:4]
            if q.get("q") and len(opts) == 4 and int(q.get("answer", -1)) in range(4):
                qs.append({"q": q["q"], "options": opts, "answer": int(q["answer"])})
        return qs[:count]

    return list(await asyncio.gather(*[one(i) for i in range(1, variants + 1)]))


def _cell_shade(cell, color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def _set_columns(section, n: int) -> None:
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "567")


def _run(p, text, *, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _header(doc, ts, L, vi, subject, topic, grade):
    kind, hc = ts["kind"], ts["head"] or "000000"
    title = f"{L['variant'].format(n=vi)}  •  {L['class'].format(n=grade)}"
    if kind == "grid":
        p = doc.add_paragraph()
        _run(p, "  " + title, bold=True, size=15, color="FFFFFF")
        _shade(p, hc)
        doc.add_paragraph(f"{L['subject']}: {subject}    {L['topic']}: {topic}")
        doc.add_paragraph(L["name"])
    elif kind in ("cards", "table"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, title, bold=True, size=16, color=hc)
        t = doc.add_table(rows=2, cols=2)
        t.style = "Table Grid"
        t.cell(0, 0).text = f"{L['subject']}: {subject}"
        t.cell(0, 1).text = f"{L['topic']}: {topic}"
        t.cell(1, 0).text = L["name"]
        t.cell(1, 1).text = L["date"]
        for c in (t.cell(0, 0), t.cell(0, 1)):
            _cell_shade(c, "F2F2F2")
        doc.add_paragraph()
    elif kind == "columns":
        p = doc.add_paragraph()
        _run(p, title, bold=True, size=15, color=hc)
        _pborder(p, ("left",), hc, 36, 8)
        doc.add_paragraph(f"{subject} — {topic}")
        doc.add_paragraph(L["name"])
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, title, bold=True, size=15)
        doc.add_paragraph(f"{L['subject']}: {subject}    {L['topic']}: {topic}")
        doc.add_paragraph(L["name"])


def _questions(doc, ts, L, qs):
    kind, hc = ts["kind"], ts["head"]
    if kind == "table":
        t = doc.add_table(rows=len(qs) + 1, cols=6)
        t.style = "Table Grid"
        for c, h in enumerate(["№", L["question"], "A", "B", "C", "D"]):
            t.cell(0, c).text = h
            _cell_shade(t.cell(0, c), hc)
            for r in t.cell(0, c).paragraphs[0].runs:
                r.bold, r.font.color.rgb = True, RGBColor.from_string("FFFFFF")
        for qi, q in enumerate(qs, 1):
            t.cell(qi, 0).text = str(qi)
            t.cell(qi, 1).text = q["q"]
            for k, o in enumerate(q["options"]):
                t.cell(qi, 2 + k).text = o
            if qi % 2 == 0:
                for c in range(6):
                    _cell_shade(t.cell(qi, c), "F7F2F4")
        widths = [Cm(1), Cm(7.4), Cm(2.15), Cm(2.15), Cm(2.15), Cm(2.15)]
        t.autofit = False
        for c, w in enumerate(widths):
            t.columns[c].width = w
        for row in t.rows:
            for c, w in enumerate(widths):
                row.cells[c].width = w
        return
    for qi, q in enumerate(qs, 1):
        if kind == "cards":
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            cell = t.cell(0, 0)
            p = cell.paragraphs[0]
            _run(p, f"{qi}. ", bold=True, color=hc)
            _run(p, q["q"], bold=True)
            cell.add_paragraph("     ".join(f"{LETTERS[i]}) {o}" for i, o in enumerate(q["options"])))
            ap = cell.add_paragraph()
            ap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(ap, f"{L['answer']}: ☐", color=hc)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, f"{qi}. {q['q']}", bold=True, color=hc if kind == "grid" else None)
        if kind == "grid":
            t = doc.add_table(rows=2, cols=2)
            for i, o in enumerate(q["options"]):
                t.cell(i // 2, i % 2).text = f"{LETTERS[i]}) {o}"
        else:
            for i, o in enumerate(q["options"]):
                op = doc.add_paragraph(f"{LETTERS[i]}) {o}")
                op.paragraph_format.left_indent = Cm(0.8)


def build_docx(variants: list[list[dict]], out: Path, subject: str, topic: str, grade: str, lang: str,
               style: int = 0) -> Path:
    L = LABELS[lang]
    ts = STYLES[style % len(STYLES)]
    hc = ts["head"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = ts["font"], Pt(11 if ts["kind"] in ("columns", "table") else 12)
    st.paragraph_format.space_after = Pt(2)
    if ts["kind"] == "columns":
        _set_columns(sec, 2)

    for vi, qs in enumerate(variants, 1):
        if vi > 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _header(doc, ts, L, vi, subject, topic, grade)
        _questions(doc, ts, L, qs)

    # Javoblar kaliti — har doim bitta ustunli yangi sahifada
    key_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if ts["kind"] == "columns":
        _set_columns(key_sec, 1)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(h, L["keys"], bold=True, color=hc)
    for vi, qs in enumerate(variants, 1):
        doc.add_paragraph().add_run(L["variant"].format(n=vi)).bold = True
        cols = 10
        rows = (len(qs) + cols - 1) // cols
        t = doc.add_table(rows=rows * 2, cols=cols)
        t.style = "Table Grid"
        for i, q in enumerate(qs):
            rr, cc = (i // cols) * 2, i % cols
            t.cell(rr, cc).text = str(i + 1)
            t.cell(rr + 1, cc).text = LETTERS[q["answer"]]
            if hc:
                _cell_shade(t.cell(rr, cc), "F2F2F2")
    doc.save(out)
    return out
