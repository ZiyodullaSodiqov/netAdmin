"""Mustaqil ish, referat, maqola (DOCX).

Avval reja (bo'limlar) tuziladi, keyin har bo'lim alohida va parallel yoziladi —
shunda kerakli hajm (varoqlar soni) chiqadi.
"""
from __future__ import annotations

import asyncio
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..ai.clients import openai_json
from ..services import LANG_FULL

WORDS_PER_PAGE = 300  # Times New Roman 14, 1.5 interval

# 5 uslub: shrift, sarlavha rangi/joylashuvi, sahifa ramkasi
STYLES = [
    {"name": "Klassik", "font": "Times New Roman", "size": 14, "spacing": 1.5, "head": None, "align": "center",
     "border": None},
    {"name": "Ramkali", "font": "Times New Roman", "size": 14, "spacing": 1.5, "head": "1F3864", "align": "center",
     "border": ("double", 6, "1F3864")},
    {"name": "Zamonaviy", "font": "Arial", "size": 12, "spacing": 1.5, "head": "0F7B6C", "align": "left",
     "border": None},
    {"name": "Rasmiy", "font": "Cambria", "size": 13, "spacing": 1.5, "head": "7A1F3D", "align": "center",
     "border": ("thickThinSmallGap", 12, "7A1F3D")},
    {"name": "Minimal", "font": "Calibri", "size": 13, "spacing": 1.3, "head": "2F5597", "align": "left",
     "border": ("single", 4, "8EA9DB")},
]

LABELS = {
    "uz": {"plan": "REJA", "intro": "Kirish", "concl": "Xulosa", "refs": "Foydalanilgan adabiyotlar",
           "by": "Bajardi", "checked": "Tekshirdi", "topic": "Mavzu", "abstract": "Annotatsiya",
           "keywords": "Kalit so'zlar", "city": "Toshkent"},
    "ru": {"plan": "ПЛАН", "intro": "Введение", "concl": "Заключение", "refs": "Список использованной литературы",
           "by": "Выполнил(а)", "checked": "Проверил(а)", "topic": "Тема", "abstract": "Аннотация",
           "keywords": "Ключевые слова", "city": "Ташкент"},
}
TYPE_TITLE = {
    "uz": {"mustaqil ish": "MUSTAQIL ISH", "referat": "REFERAT", "ilmiy maqola": ""},
    "ru": {"mustaqil ish": "САМОСТОЯТЕЛЬНАЯ РАБОТА", "referat": "РЕФЕРАТ", "ilmiy maqola": ""},
}


async def generate_content(doc_type: str, topic: str, pages: int, lang: str, details: str = "") -> dict:
    L = LABELS[lang]
    system = f"Sen akademik matnlar yozuvchi mutaxassissan. Faqat JSON qaytar. Til: {LANG_FULL[lang]}."
    n_main = max(2, min(6, pages // 3))
    plan = await openai_json(system, (
        f"Hujjat turi: {doc_type}. Mavzu: {topic}. {('Talablar: ' + details) if details else ''}\n"
        f"{n_main} ta asosiy bo'lim nomini tuz (kirish va xulosadan tashqari). "
        'Format: {"title": "aniq sarlavha", "sections": ["..."], "keywords": ["5-7 ta"], '
        '"references": ["8-10 ta real, mavjud adabiyot: muallif, nom, nashriyot, yil"]}'
    ))
    sections = [s for s in plan.get("sections", []) if s][:n_main] or [topic]
    total_words = pages * WORDS_PER_PAGE
    # Titul va reja ~1.5 varoq oladi
    body_words = max(600, total_words - int(1.5 * WORDS_PER_PAGE))
    intro_w = concl_w = int(body_words * 0.1)
    sec_w = int((body_words - intro_w - concl_w) / len(sections))

    async def write(heading: str, words: int, role: str) -> dict:
        d = await openai_json(system, (
            f"{doc_type} — mavzu: «{plan.get('title', topic)}». Umumiy reja: {', '.join(sections)}.\n"
            f"Hozir faqat «{heading}» qismini yoz ({role}). Hajmi taxminan {words} so'z, "
            "4-8 ta abzats, ilmiy uslubda, aniq misollar va faktlar bilan, takrorsiz.\n"
            'Format: {"paragraphs": ["..."]}'
        ))
        return {"heading": heading, "paragraphs": [p for p in d.get("paragraphs", []) if p]}

    parts = await asyncio.gather(
        write(L["intro"], intro_w, "kirish"),
        *[write(s, sec_w, f"{i}-asosiy bo'lim") for i, s in enumerate(sections, 1)],
        write(L["concl"], concl_w, "xulosa"),
    )
    result = {"title": plan.get("title", topic), "sections": list(parts),
              "keywords": plan.get("keywords", []), "references": plan.get("references", [])}
    if doc_type == "ilmiy maqola":
        ab = await openai_json(system, (
            f"«{result['title']}» maqolasi uchun 60-90 so'zli annotatsiya yoz. "
            'Format: {"abstract": "...", "abstract_en": "inglizcha tarjimasi"}'))
        result.update(ab)
    return result


def _page_border(sec, kind: str, size: int, color: str) -> None:
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), kind)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), color)
        borders.append(el)
    # OOXML sxemasida pgBorders pgMar dan keyin turishi shart (aks holda Word faylni "buzilgan" deydi)
    pg_mar = sec._sectPr.find(qn("w:pgMar"))
    if pg_mar is not None:
        pg_mar.addnext(borders)
    else:
        sec._sectPr.append(borders)


def _base_doc(st: dict | None = None) -> Document:
    st = st or STYLES[0]
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = st["font"], Pt(st["size"])
    normal.paragraph_format.line_spacing = st["spacing"]
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2"):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold = st["font"], Pt(st["size"] + (2 if name == "Heading 1" else 0)), True
        h.font.color.rgb = RGBColor.from_string(st["head"]) if st["head"] else None
    if st["border"]:
        _page_border(sec, *st["border"])
    return doc


def _para(doc, text, *, bold=False, size=None, align=None, indent=True, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        p.alignment = align
    elif indent:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_docx(data: dict, out: Path, doc_type: str, lang: str, author: str = "", institution: str = "",
               teacher: str = "", style: int = 0) -> Path:
    L = LABELS[lang]
    st = STYLES[style % len(STYLES)]
    doc = _base_doc(st)
    C = WD_ALIGN_PARAGRAPH.CENTER
    HC = st["head"]  # sarlavha rangi
    HA = C if st["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT

    if doc_type == "ilmiy maqola":
        _para(doc, data["title"].upper(), bold=True, size=16, align=C, color=HC)
        if author:
            _para(doc, author, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if institution:
            _para(doc, institution, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if data.get("abstract"):
            _para(doc, f"{L['abstract']}: ", bold=True, indent=False)
            _para(doc, data["abstract"])
        if data.get("keywords"):
            p = _para(doc, f"{L['keywords']}: ", bold=True, indent=False)
            p.add_run(", ".join(data["keywords"])).italic = True
        if data.get("abstract_en") and lang != "en":
            _para(doc, "Abstract: ", bold=True, indent=False)
            _para(doc, data["abstract_en"])
        for s in data["sections"]:
            _para(doc, s["heading"], bold=True, align=HA, color=HC)
            for t in s["paragraphs"]:
                _para(doc, t)
    else:
        # Titul varaq
        if institution:
            _para(doc, institution.upper(), bold=True, align=C)
        for _ in range(6):
            doc.add_paragraph()
        _para(doc, TYPE_TITLE[lang][doc_type], bold=True, size=28, align=C, color=HC)
        _para(doc, f"{L['topic']}: {data['title']}", bold=True, size=16, align=C)
        for _ in range(6):
            doc.add_paragraph()
        if author:
            _para(doc, f"{L['by']}: {author}", align=WD_ALIGN_PARAGRAPH.RIGHT)
        if teacher:
            _para(doc, f"{L['checked']}: {teacher}", align=WD_ALIGN_PARAGRAPH.RIGHT)
        for _ in range(3):
            doc.add_paragraph()
        _para(doc, L["city"], align=C)
        _page_break(doc)
        # Reja
        _para(doc, L["plan"], bold=True, size=16, align=C, color=HC)
        for i, s in enumerate(data["sections"]):
            label = s["heading"] if i in (0, len(data["sections"]) - 1) else f"{i}. {s['heading']}"
            _para(doc, label, indent=False)
        _para(doc, L["refs"], indent=False)
        _page_break(doc)
        for s in data["sections"]:
            doc.add_heading(s["heading"], level=1).alignment = HA
            for t in s["paragraphs"]:
                _para(doc, t)

    if data.get("references"):
        if doc_type != "ilmiy maqola":
            _page_break(doc)
        _para(doc, L["refs"], bold=True, align=C, color=HC)
        for i, ref in enumerate(data["references"], 1):
            _para(doc, f"{i}. {ref}", indent=False)

    _add_page_numbers(doc)
    doc.save(out)
    return out


def _add_page_numbers(doc) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)

