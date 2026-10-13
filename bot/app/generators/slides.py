"""Taqdimot (PPTX) — matn OpenAI dan, ko'rinish python-pptx bilan."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from ..ai.clients import openai_json
from ..services import LANG_FULL

# 5 uslub: rang, shrift va joylashuv (top — yuqorida sarlavha paneli, side — chap tomonda chiziq)
THEMES = [
    {"bg": "FFFFFF", "accent": "1F4E79", "accent2": "2E75B6", "text": "222222", "title_fg": "FFFFFF",
     "head": "FFFFFF", "font": "Calibri", "layout": "top", "bullet": "▪"},
    {"bg": "FBFAF7", "accent": "7A1F3D", "accent2": "C9A227", "text": "2B2B2B", "title_fg": "FFFFFF",
     "head": "7A1F3D", "font": "Georgia", "layout": "side", "bullet": "◆"},
    {"bg": "F4FAF6", "accent": "1E6B52", "accent2": "3BA57A", "text": "1F2A24", "title_fg": "FFFFFF",
     "head": "FFFFFF", "font": "Arial", "layout": "top", "bullet": "●"},
    {"bg": "1E2230", "accent": "F5A524", "accent2": "4FC3F7", "text": "E8EAF0", "title_fg": "1E2230",
     "head": "F5A524", "font": "Trebuchet MS", "layout": "side", "bullet": "➤"},
    {"bg": "FFFFFF", "accent": "5B3FA8", "accent2": "18A999", "text": "2A2540", "title_fg": "FFFFFF",
     "head": "5B3FA8", "font": "Cambria", "layout": "line", "bullet": "—"},
]

END_TEXT = {"uz": "E'tiboringiz uchun rahmat!", "ru": "Спасибо за внимание!", "en": "Thank you for your attention!"}


async def generate_content(topic: str, slides: int, lang: str, details: str = "") -> dict:
    system = (
        "Sen tajribali o'qituvchi va taqdimot dizaynerisan. Faqat JSON qaytar. "
        f"Barcha matn {LANG_FULL[lang]} tilida bo'lsin."
    )
    user = (
        f"Mavzu: {topic}\n{('Qo`shimcha talablar: ' + details) if details else ''}\n"
        f"Jami {slides} ta slayd: 1-slayd sarlavha, oxirgisi \"rahmat\" slaydi (ularni men qo'shaman), oxirgidan oldingisi xulosa. Har bir slaydda 3-5 ta mazmunli, "
        "aniq faktlarga asoslangan punkt (har biri 8-20 so'z). Takrorlanmasin.\n"
        'Format: {"title": "...", "subtitle": "...", "slides": [{"title": "...", "bullets": ["..."]}]} '
        f"— slides massivida aynan {slides - 2} ta element bo'lsin (sarlavha va rahmat slaydlarisiz)."
    )
    data = await openai_json(system, user)
    data.setdefault("title", topic)
    data["slides"] = [s for s in data.get("slides", []) if s.get("title")][: slides - 2]
    return data


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _rect(slide, x, y, w, h, color: str):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    return shp


def _bg(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def build_pptx(data: dict, out: Path, lang: str, author: str = "", theme: int = 0) -> Path:
    t = THEMES[theme % len(THEMES)]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    # Sarlavha slaydi
    s = prs.slides.add_slide(blank)
    _bg(s, t["accent"])
    _rect(s, 0, int(H * 0.62), W, int(H * 0.012), t["accent2"])
    tb = s.shapes.add_textbox(int(W * 0.07), int(H * 0.25), int(W * 0.86), int(H * 0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = data["title"]
    p.font.size, p.font.bold, p.font.color.rgb = Pt(40), True, _rgb(t["title_fg"])
    p.font.name = t["font"]
    if data.get("subtitle"):
        p2 = tf.add_paragraph()
        p2.text = data["subtitle"]
        p2.font.size, p2.font.color.rgb, p2.font.name = Pt(20), _rgb(t["title_fg"]), t["font"]
    if author:
        ab = s.shapes.add_textbox(int(W * 0.07), int(H * 0.7), int(W * 0.86), int(H * 0.15))
        ab.text_frame.text = author
        ab.text_frame.paragraphs[0].font.size = Pt(18)
        ab.text_frame.paragraphs[0].font.color.rgb = _rgb(t["title_fg"])

    # Mazmun slaydlari
    total = len(data["slides"])
    for i, sl in enumerate(data["slides"], start=2):
        s = prs.slides.add_slide(blank)
        _bg(s, t["bg"])
        if t["layout"] == "top":
            _rect(s, 0, 0, W, int(H * 0.16), t["accent"])
            _rect(s, 0, int(H * 0.16), int(W * 0.25), int(H * 0.01), t["accent2"])
        elif t["layout"] == "side":
            _rect(s, 0, 0, int(W * 0.025), H, t["accent"])
            _rect(s, int(W * 0.025), 0, int(W * 0.006), H, t["accent2"])
        else:  # line
            _rect(s, int(W * 0.05), int(H * 0.165), int(W * 0.9), int(H * 0.004), t["accent2"])
        hb = s.shapes.add_textbox(int(W * 0.05), int(H * 0.03), int(W * 0.9), int(H * 0.11))
        hb.text_frame.word_wrap = True
        hp = hb.text_frame.paragraphs[0]
        hp.text = sl["title"]
        hp.font.size, hp.font.bold, hp.font.color.rgb = Pt(28), True, _rgb(t["head"])
        hp.font.name = t["font"]

        bb = s.shapes.add_textbox(int(W * 0.06), int(H * 0.22), int(W * 0.88), int(H * 0.68))
        tf = bb.text_frame
        tf.word_wrap = True
        bullets = sl.get("bullets") or []
        size = 20 if sum(len(b) for b in bullets) < 420 else 17
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"{t['bullet']} " + b
            p.font.size, p.font.color.rgb, p.font.name = Pt(size), _rgb(t["text"]), t["font"]
            p.space_after = Pt(10)
        nb = s.shapes.add_textbox(int(W * 0.9), int(H * 0.92), int(W * 0.08), int(H * 0.06))
        nb.text_frame.text = f"{i}/{total + 2}"
        nb.text_frame.paragraphs[0].font.size = Pt(12)
        nb.text_frame.paragraphs[0].font.color.rgb = _rgb(t["text"])
        nb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    # Yakuniy slayd
    s = prs.slides.add_slide(blank)
    _bg(s, t["accent"])
    tb = s.shapes.add_textbox(int(W * 0.07), int(H * 0.4), int(W * 0.86), int(H * 0.2))
    p = tb.text_frame.paragraphs[0]
    p.text = END_TEXT.get(lang, END_TEXT["uz"])
    p.alignment = PP_ALIGN.CENTER
    p.font.size, p.font.bold, p.font.color.rgb = Pt(40), True, _rgb(t["title_fg"])
    p.font.name = t["font"]

    prs.save(out)
    return out
